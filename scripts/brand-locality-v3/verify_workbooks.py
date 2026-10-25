"""Read-only XLSX verification against the completed pipeline's CSV exports."""

import argparse
import csv
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


def read(path):
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def same_number(value, original, tolerance=1e-8):
    if not original:
        assert value is None
    else:
        assert isinstance(value, int | float) and not isinstance(value, bool)
        assert abs(value - float(original)) <= tolerance


def data_rows(sheet):
    headers = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    return [
        dict(zip(headers, row, strict=True)) for row in sheet.iter_rows(min_row=7, values_only=True)
    ]


def native_checks(path):
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    formulas, errors, tables = 0, [], 0
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                tables += 1
            if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
                continue
            with archive.open(name) as stream:
                for _, element in ElementTree.iterparse(stream, events=("end",)):
                    if element.tag.endswith("}c"):
                        if element.attrib.get("t") == "e":
                            errors.append(
                                (
                                    name,
                                    element.attrib.get("r"),
                                    element.findtext("s:v", namespaces=namespace),
                                )
                            )
                        if element.find("s:f", namespace) is not None:
                            formulas += 1
                        element.clear()
    assert not errors, errors[:10]
    assert formulas > 0 and tables >= 2
    return {"formulas": formulas, "tables": tables, "error_cells": 0, "zip_valid": True}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run", default="output/geocoding-v3")
    parser.add_argument("--output", default="outputs/truckroute-brand-locality-v3")
    args = parser.parse_args()
    run, output = Path(args.run), Path(args.output)
    report = json.loads((run / "report.json").read_text())
    sources = read(run / "source-comparison.csv")
    stops = read(run / "final-stops.csv")
    prices = read(run / "fuel-prices-geocoded.csv")
    analysis_file = output / "source-matching-analysis.xlsx"
    final_file = output / "final-geocoded-locations.xlsx"
    native = {str(file.name): native_checks(file) for file in (analysis_file, final_file)}
    analysis = load_workbook(analysis_file, read_only=True, data_only=True)
    assert analysis.sheetnames == ["Summary", "Source matches", "Coordinate disagreements"]
    source_rows = data_rows(analysis["Source matches"])
    assert len(source_rows) == len(sources) == 20214
    for row, original in zip(source_rows, sources, strict=True):
        assert row["OPIS stop ID"] == original["external_id"]
        assert row["Provider"] == original["provider"]
        same_number(row["Source score"], original["score"], 0.00051)
        for title, key in (
            ("Name points", "name_points"),
            ("Locality points", "locality_points"),
            ("Highway points", "highway_points"),
            ("Exit points", "exit_points"),
        ):
            # A blank input criterion is displayed as a zero contribution by Excel.
            if original["applicable_weight"]:
                same_number(row[title], original[key] or "0", 0.0000011)
            else:
                assert row[title] is None
        same_number(row["Source latitude"], original["predicted_latitude"])
        same_number(row["Source longitude"], original["predicted_longitude"])
        assert row["Eligible point"] == (original["eligible_point"] == "True")
    for row_number, provider in enumerate(("alltheplaces", "overture", "osm"), start=6):
        assert (
            analysis["Summary"][f"C{row_number}"].value
            == report["source_candidate_coordinates"][provider]
        )
    assert (
        len(data_rows(analysis["Coordinate disagreements"]))
        == report["coordinate_deviation"]["flagged_at_1km"]
    )
    analysis.close()
    final = load_workbook(final_file, read_only=True, data_only=True)
    assert final.sheetnames == ["Summary", "Final stops", "Fuel prices"]
    final_rows = data_rows(final["Final stops"])
    assert len(final_rows) == len(stops) == 6738
    for row, original in zip(final_rows, stops, strict=True):
        assert row["OPIS stop ID"] == original["external_id"]
        assert row["Status"] == original["status"]
        assert row["Location method"] == original["prediction_method"]
        assert row["Approximate"] == (original["is_approximate"] == "True")
        for label, key in (
            ("Predicted latitude", "predicted_latitude"),
            ("Predicted longitude", "predicted_longitude"),
            ("Accepted latitude", "latitude"),
            ("Accepted longitude", "longitude"),
        ):
            same_number(row[label], original[key])
    price_rows = data_rows(final["Fuel prices"])
    assert len(price_rows) == len(prices) == 8151
    for row, original in zip(price_rows, prices, strict=True):
        for label, key in (
            ("OPIS stop ID", "OPIS Truckstop ID"),
            ("Truckstop name", "Truckstop Name"),
            ("Original address", "Address"),
            ("City", "City"),
            ("State", "State"),
            ("Rack ID", "Rack ID"),
        ):
            assert row[label] == original[key]
        same_number(row["Retail price"], original["Retail Price"], 0.0000000051)
    for index, status in enumerate(
        ("accepted", "needs_review", "approximate", "unmatched"), start=6
    ):
        assert final["Summary"][f"B{index}"].value == report["geocodes"][status]
    final.close()
    result = {
        "passed": True,
        "source_rows": len(sources),
        "final_stops": len(stops),
        "original_price_rows": len(prices),
        "native_xlsx_checks": native,
    }
    (run / "workbook-validation.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
